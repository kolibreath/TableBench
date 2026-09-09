from __future__ import annotations

import datetime
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time

if getattr(sys, 'frozen', False):
    sys.path.insert(0, sys._MEIPASS)

import requests
from flask import Flask, jsonify, request, Response, stream_with_context

from ai_provider import MockAIProvider
from parser import (extract_all_subsections_with_content,
                    extract_text_from_wps, parse_toc_structure,
                    check_toc_body_consistency)
import parser
import version_manager
import doc_engine
import storage_service
import workhours_service

# ── 自动清理配置 ─────────────────────────────────────────────────
CLEANUP_INTERVAL_SECONDS = 300       # 清理间隔：5 分钟
CLEANUP_EXPIRY_SECONDS = 1800        # 过期时间：30 分钟


def _get_temp_dir() -> str:
    """获取临时目录，优先使用 exe 所在路径，避免内网环境系统临时目录权限问题"""
    if getattr(sys, 'frozen', False):
        temp_dir = os.path.join(os.path.dirname(sys.executable), "temp_uploads")
    else:
        temp_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "temp_uploads")

    os.makedirs(temp_dir, exist_ok=True)
    return temp_dir


def _get_temp_session_dir() -> str:
    """以当前请求时间戳创建临时会话目录"""
    base_dir = _get_temp_dir()
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    session_dir = os.path.join(base_dir, timestamp)
    os.makedirs(session_dir, exist_ok=True)
    return session_dir


app = Flask(__name__)

_batch_docs: dict[str, dict] = {}

# ── Mock 模式下的 session 元数据（用于 agentInitChat 生成对应数量结果）─
_mock_session_meta: dict[str, dict] = {}

# ── AI 模式检测 ────────────────────────────────────────────────
_IS_FROZEN = getattr(sys, "frozen", False)
_MOCK_MODE_ENV = os.environ.get("AI_MOCK_MODE", "").lower()

if _IS_FROZEN:
    AI_MOCK_MODE = _MOCK_MODE_ENV == "true"
    print(f"[AI] 检测到 PyInstaller 环境，默认真实模式（可 AI_MOCK_MODE=true 强制 Mock）")
else:
    AI_MOCK_MODE = _MOCK_MODE_ENV != "false"
    print(f"[AI] 检测到源码环境，默认 Mock 模式（可 AI_MOCK_MODE=false 强制真实）")

print(f"[AI] 当前 AI 模式: {'Mock' if AI_MOCK_MODE else '真实'}")


# ── AI 大模型接口配置 ──────────────────────────────────────────
AI_BASE_URLS_CONFIG_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "AI_BASE_URLS.json",
)

_DEFAULT_AI_URLS: list[str] = []


def load_ai_base_urls() -> list[str]:
    if os.path.exists(AI_BASE_URLS_CONFIG_FILE):
        try:
            with open(AI_BASE_URLS_CONFIG_FILE, "r", encoding="utf-8") as f:
                config = json.load(f)
            urls = config.get("urls", [])
            if isinstance(urls, list) and all(isinstance(u, str) for u in urls):
                print(f"[AI] 从配置文件加载接口列表: {len(urls)} 个")
                return urls
        except Exception as e:
            print(f"[AI] 读取配置文件失败: {e}")
    print("[AI] 未找到配置文件，使用默认接口列表（空）")
    return _DEFAULT_AI_URLS


AI_BASE_URLS = load_ai_base_urls()


def check_ai_interface_health(base_url: str) -> bool:
    try:
        resp = requests.get(f"{base_url}/chatabc/health_check", timeout=5)
        return resp.status_code == 200
    except Exception as e:
        print(f"[AI] 健康检查失败 ({base_url}): {e}")
        return False


def get_available_ai_base_url() -> str | None:
    if not AI_BASE_URLS:
        print("[AI] 接口列表为空")
        return None
    for url in AI_BASE_URLS:
        print(f"[AI] 检查接口: {url}")
        if check_ai_interface_health(url):
            print(f"[AI] ✓ 接口可用: {url}")
            return url
        print(f"[AI] ✗ 接口不可用: {url}")
    print("[AI] 所有接口均不可用")
    return None


def ai_init_session(base_url: str) -> str | None:
    payload = {
        "appId": "", "trCode": "", "trVersion": "",
        "timestamp": "1", "requestId": "",
        "data": {"prompt_variables": [{"name": "", "value": ""}]},
    }

    # 打印发送给 AI 大模型的完整报文
    print("\n" + "=" * 80)
    print(f"[AI] 发送给 AI 大模型的 init_session 报文 (baseUrl: {base_url}):")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    print("=" * 80 + "\n")

    try:
        resp = requests.post(f"{base_url}/chatabc/init_session", json=payload, timeout=30)
        if resp.status_code != 200:
            print(f"[AI] init_session 失败: {resp.status_code}")
            return None
        data = resp.json()
        session_id = data.get("data", {}).get("session_id")
        if session_id:
            print(f"[AI] session_id 获取成功: {session_id}")
            return session_id
        print("[AI] init_session 返回格式错误")
        return None
    except Exception as e:
        print(f"[AI] init_session 异常: {e}")
        return None


def ai_chat(base_url: str, session_id: str, prompt: str) -> str | None:
    payload = {
        "appId": "", "trCode": "", "trVersion": "",
        "timestamp": "1", "requestId": "",
        "data": {
            "session_id": session_id,
            "txt": prompt,
            "files": [],
            "stream": False,
        },
    }

    # 打印发送给 AI 大模型的完整报文
    print("\n" + "=" * 80)
    print(f"[AI] 发送给 AI 大模型的 chat 报文 (session_id: {session_id}):")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    print("=" * 80 + "\n")

    try:
        resp = requests.post(f"{base_url}/chatabc/chat", json=payload, timeout=120)
        if resp.status_code != 200:
            print(f"[AI] chat 请求失败: {resp.status_code}")
            return None
        body = resp.json()
        json_res = body.get("data", {}).get("json_res", "")
        if json_res:
            return json_res
        return body.get("data", {}).get("txt", "")
    except Exception as e:
        print(f"[AI] chat 异常: {e}")
        return None


# ── CORS ──────────────────────────────────────────────────────

@app.after_request
def add_cors_headers(response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type"
    response.headers["Access-Control-Allow-Methods"] = "POST, OPTIONS"
    return response



# ═══════════════════════════════════════════════════════════
#  V3.0 统一后端：健康检查 / 文档引擎 / 检查历史 / 数据收集 / 文档自查
# ═══════════════════════════════════════════════════════════

@app.route("/health", methods=["GET"])
def health():
    """健康检查 + 能力位（popup / check / estimation 页探活用）"""
    engine = doc_engine.detect_engine()
    capabilities = {"parseCom": engine == "wps_com", "docCheck": True}
    version = "3.0.0"
    try:
        with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "version.json"), "r", encoding="utf-8") as f:
            version = json.load(f).get("version", version)
    except Exception:
        pass
    return jsonify({
        "ok": True,
        "service": "filechecker-backend",
        "version": version,
        "docEngine": engine,
        "capabilities": capabilities,
    })


@app.route("/api/parse_doc_com", methods=["POST"])
def api_parse_doc_com():
    """WPS COM 标题树解析（pyparser 环境自动降级）
    入参二选一：multipart file=文档 ｜ JSON {files:[{name,data(base64)}]} 取第一个
    """
    temp_path = None
    try:
        if "file" in request.files:
            file = request.files["file"]
            safe_name = "".join(c if c.isalnum() or c in "._- " else "_" for c in file.filename)
            temp_dir = _get_temp_session_dir()
            temp_path = os.path.join(temp_dir, f"parse_{safe_name}")
            file.save(temp_path)
        else:
            body = request.get_json(silent=True) or {}
            files = body.get("files") or []
            if not files:
                return jsonify({"error": "缺少文件"}), 400
            import base64 as _b64
            temp_dir = _get_temp_session_dir()
            temp_path = os.path.join(temp_dir, f"parse_{files[0].get('name', 'doc')}")
            with open(temp_path, "wb") as fh:
                fh.write(_b64.b64decode(files[0].get("data", "")))

        result = doc_engine.extract(temp_path)
        # 控制返回体积：全文裁剪到 30 万字符
        if len(result.get("fulltext", "")) > 300000:
            result["fulltext"] = result["fulltext"][:300000] + "…[截断]"
        return jsonify(result)
    except Exception as e:
        print(f"[parse_doc_com] 解析失败: {e}")
        return jsonify({"error": f"解析失败: {e}"}), 500
    finally:
        if temp_path:
            try:
                os.remove(temp_path)
            except Exception:
                pass


# ── 检查历史（IndexedDB 主存的 SQLite 备份） ──────────────────

@app.route("/api/history/sync", methods=["POST"])
def api_history_sync():
    record = request.get_json(silent=True) or {}
    return jsonify(storage_service.history_sync(record))


@app.route("/api/history/list", methods=["POST", "GET"])
def api_history_list():
    filter_ = request.get_json(silent=True) or {}
    return jsonify(storage_service.history_list(filter_))


@app.route("/api/history/get", methods=["POST"])
def api_history_get():
    body = request.get_json(silent=True) or {}
    return jsonify(storage_service.history_get(body.get("id")))


@app.route("/api/history/delete", methods=["POST"])
def api_history_delete():
    body = request.get_json(silent=True) or {}
    return jsonify(storage_service.history_delete(body.get("id")))


# ── 合规检查数据收集（JSONL 按月落盘） ────────────────────────

@app.route("/api/collect", methods=["POST"])
def api_collect():
    payload = request.get_json(silent=True) or {}
    return jsonify(storage_service.collect_append(payload))


@app.route("/api/collect/export", methods=["GET"])
def api_collect_export():
    month = request.args.get("month", "")
    try:
        zip_path = storage_service.collect_export(month or None)
        if not zip_path:
            return jsonify({"error": "没有可导出的收集数据"}), 404
        from flask import send_file
        return send_file(zip_path, as_attachment=True)
    except Exception as e:
        return jsonify({"error": f"导出失败: {e}"}), 500


# ── 项目文档自查（原 doc_check 本地服务迁入） ─────────────────

@app.route("/api/check_upload", methods=["POST"])
def api_doc_check_upload():
    """接收浏览器上传的文档(base64)，写临时目录后用文档引擎检查"""
    import base64 as _b64
    body = request.get_json(silent=True) or {}
    files = body.get("files") or []
    if not files:
        return jsonify({"error": "缺少要检查的文件"}), 400

    import tempfile as _tempfile
    tmpdir = _tempfile.mkdtemp(prefix="doc_check_")
    paths = []
    try:
        for f in files:
            name = f.get("name", "upload.bin")
            p = os.path.join(tmpdir, name.replace("/", "_"))
            try:
                with open(p, "wb") as fh:
                    fh.write(_b64.b64decode(f.get("data", "")))
                paths.append(p)
            except Exception:
                continue

        # 优先 WPS COM 引擎（check_docs），不可用时给出明确错误
        if doc_engine.detect_engine() != "wps_com":
            return jsonify({
                "error": "文档自查需要 Windows + WPS 环境（WPS COM）。"
                         "当前为纯 Python 引擎（docEngine=pyparser），估算书合规检查不受影响。",
                "documents": [],
            })
        import check_docs as C
        return jsonify(C.check_files(paths))
    finally:
        import shutil as _shutil
        _shutil.rmtree(tmpdir, ignore_errors=True)


# ── 工时填报检查（V3.1 并入，原 WorkingHoursTool app.exe 逻辑） ──

@app.route("/api/workhours/projects", methods=["POST"])
def api_workhours_projects():
    """有工时计划的项目列表（今年+去年合并）。body: {cookie}"""
    body = request.get_json(silent=True) or {}
    cookie = (body.get("cookie") or "").strip()
    if not cookie:
        return jsonify({"ok": False, "error": "缺少 ITA Cookie，请先登录 ita.abc"}), 400
    try:
        return jsonify(workhours_service.get_projects_two_year(cookie))
    except Exception as e:
        return jsonify({"ok": False, "error": f"获取工时项目失败：{e}"}), 502


@app.route("/api/workhours/target", methods=["POST"])
def api_workhours_target():
    """项目成员工时填报统计（今年+去年合并）。body: {cookie, prjid}"""
    body = request.get_json(silent=True) or {}
    cookie = (body.get("cookie") or "").strip()
    prjid = (body.get("prjid") or "").strip()
    if not cookie or not prjid:
        return jsonify({"ok": False, "error": "缺少 cookie 或 prjid"}), 400
    try:
        return jsonify(workhours_service.get_target_two_year(cookie, prjid))
    except Exception as e:
        return jsonify({"ok": False, "error": f"获取成员工时失败：{e}"}), 502


# ── 文档上传 / 匹配 / 重置 ──────────────────────────────────

@app.route("/api/upload", methods=["POST"])
def upload():
    if "file" not in request.files:
        return jsonify({"error": "缺少文件"}), 400
    batch_name = request.form.get("batchName", "")
    if not batch_name:
        return jsonify({"error": "缺少 batchName"}), 400

    file = request.files["file"]
    safe_name = "".join(c if c.isalnum() or c in "._- " else "_" for c in file.filename)
    temp_dir = _get_temp_session_dir()
    temp_path = os.path.join(temp_dir, f"req_{batch_name}_{safe_name}")
    file.save(temp_path)

    try:
        text = extract_text_from_wps(temp_path)
    except Exception as e:
        try:
            os.remove(temp_path)
        except Exception:
            pass
        return jsonify({"error": f"解析文件失败: {e}"}), 400

    try:
        os.remove(temp_path)
    except PermissionError:
        pass

    toc = parse_toc_structure(text)
    content_map = extract_all_subsections_with_content(text, toc)
    _batch_docs[batch_name] = {"toc": toc, "content_map": content_map}

    print(f"\n{'=' * 60}")
    print(f"【上传批次: {batch_name}】")
    print(f"{'=' * 60}")
    print("【TOC 目录结构】")
    for key in sorted(toc.keys(), key=lambda k: (int(k.split(".")[0]), float(k) if "." in k else 0)):
        print(f"  [{key}] {toc[key]}")
    print(f"\n【功能需求子节正文 Map】")
    for key in sorted(content_map.keys(), key=lambda k: float(k)):
        item = content_map[key]
        preview = item["content"][:100] if item["content"] else "(无内容)"
        print(f"  [{key}] {item['title']}")
        print(f"      内容预览: {preview}")

    return jsonify({
        "status": "ok",
        "batchName": batch_name,
        "toc_count": len(toc),
        "content_count": len(content_map),
    })


@app.route("/api/match", methods=["POST"])
def match():
    data = request.get_json(force=True)
    batch_name = data.get("batchName", "")
    chapter = data.get("chapter", "")
    keywords = data.get("keywords", [])

    if not batch_name or not keywords:
        return jsonify({"error": "缺少必要参数: batchName, keywords"}), 400

    if batch_name not in _batch_docs:
        return jsonify({"error": f"批次 {batch_name} 未上传"}), 400

    content_map = _batch_docs[batch_name]["content_map"]
    if chapter and chapter in content_map:
        subsection_content = content_map[chapter].get("content", "")
    else:
        subsection_content = ""

    found = [kw for kw in keywords if kw in subsection_content]
    not_found = [kw for kw in keywords if kw not in subsection_content]

    return jsonify({
        "found": found,
        "not_found": not_found,
        "all_found": len(not_found) == 0,
    })


@app.route("/api/reset", methods=["POST"])
def reset():
    _batch_docs.clear()
    _mock_session_meta.clear()
    print("\n【重置】所有缓存数据已清除")
    return jsonify({"status": "ok"})


def _do_cleanup_temp_files() -> int:
    """核心清理逻辑：删除 temp_uploads/ 下超过 CLEANUP_EXPIRY_SECONDS 的会话目录。

    Returns:
        已删除的文件夹数量
    """
    base_dir = _get_temp_dir()
    now = datetime.datetime.now()
    deleted_count = 0
    if os.path.exists(base_dir):
        for folder_name in os.listdir(base_dir):
            folder_path = os.path.join(base_dir, folder_name)
            if not os.path.isdir(folder_path):
                continue
            try:
                folder_time = datetime.datetime.strptime(folder_name, "%Y%m%d_%H%M%S")
                if (now - folder_time).total_seconds() > CLEANUP_EXPIRY_SECONDS:
                    shutil.rmtree(folder_path)
                    deleted_count += 1
                    print(f"[自动清理] 已删除过期临时文件夹: {folder_name}")
            except ValueError:
                continue
    if deleted_count > 0:
        print(f"[自动清理] 完成，共删除 {deleted_count} 个过期文件夹")
    return deleted_count


@app.route("/api/cleanup-temp", methods=["POST"])
def cleanup_temp():
    """清理 temp_uploads 中超过 30 分钟的临时文件（手动触发）"""
    deleted_count = _do_cleanup_temp_files()
    return jsonify({"status": "ok", "deleted": deleted_count})


# ── AI 大模型接口 ────────────────────────────────────────────

@app.route("/api/ai/check-available", methods=["POST"])
def ai_check_available():
    data = request.get_json(silent=True) or {}
    url = data.get("url", "").strip()

    if AI_MOCK_MODE:
        print(f"[AI] Mock 模式: check-available -> True")
        return jsonify({"available": True, "url": url or "mock://ai"})

    if url:
        ok = check_ai_interface_health(url)
        print(f"[AI] 外部传入 URL 可用性检查: {url} -> {ok}")
        return jsonify({"available": ok, "url": url})
    available = get_available_ai_base_url() is not None
    return jsonify({"available": available})


@app.route("/api/ai/check-url", methods=["POST"])
def ai_check_url():
    data = request.get_json(silent=True) or {}
    url = data.get("url", "").strip()
    if not url:
        return jsonify({"ok": False, "error": "缺少 url"}), 400

    if AI_MOCK_MODE:
        session_id = MockAIProvider.init_session(url)
        print(f"[AI] Mock 模式: check-url -> session_id={session_id}")
        return jsonify({"ok": True, "session_id": session_id, "url": url})

    health = check_ai_interface_health(url)
    if not health:
        print(f"[AI] 测试连接失败（健康检查）: {url}")
        return jsonify({"ok": False, "error": "健康检查失败，请确认 AI 服务已启动"})
    session_id = ai_init_session(url)
    if not session_id:
        print(f"[AI] 测试连接失败（init_session）: {url}")
        return jsonify({"ok": False, "error": "init_session 失败，请检查 AI 接口是否正确"})
    print(f"[AI] 测试连接成功: {url}")
    return jsonify({"ok": True, "session_id": session_id, "url": url})


def _build_rule8_prompt(items: list[dict]) -> str:
    """构造规则8（重复功能名称校验）的 AI prompt"""
    prompt = (
        "[RULE_TYPE:8]\n"
        "你是一个功能点估算（Function Point Estimation）审核专家助手。"
        "请判断以下每一对功能名称是否为实质性的重复或写法相近。\n\n"
        "判断标准：\n"
        "1. 两个功能名称描述的是同一个业务操作（如\"新增用户\"和\"添加用户\"），"
        "应判定为重复（aiConfirmed=true）。\n"
        "2. 核心动词和核心宾语相同，仅修饰词不同（如\"用户信息修改\"和\"修改用户信息\"），"
        "应判定为相近（aiConfirmed=true）。\n"
        "3. 如果两个功能名称虽然用词相似但指代完全不同的业务对象或操作"
        "（如\"新增用户\"和\"新增用户组\"），则不应判定为重复（aiConfirmed=false）。\n"
        "4. 注意区分ILF内部逻辑文件（通常以\"表\"/\"文件\"结尾）和事务功能（EI/EQ/EO）。"
        "两者即使名称相似也属于不同类型，不应判定为重复。\n\n"
        "请按以下JSON格式返回结果，不要包含其他内容：\n"
        '[\n'
        '  {"nameI": "名称1", "nameJ": "名称2", "aiConfirmed": true/false, "reason": "判断理由"}\n'
        ']\n\n'
        "检查项列表：\n"
    )
    for idx, it in enumerate(items, 1):
        kind_label = "完全重复" if it.get("rule8Kind") == "dup" else "写法相近"
        prompt += (
            f"\n{idx}. 功能名称A: {it['nameI']}\n"
            f"   功能名称B: {it['nameJ']}\n"
            f"   引擎判定: {kind_label}\n"
            f"   所在Sheet: {it.get('context', {}).get('sheet', '')}\n"
            f"   分组键(A-E列): {it.get('context', {}).get('aeGroup', '')}\n"
        )
    return prompt


def _build_rule14_prompt(items: list[dict]) -> str:
    """构造规则14（单名词差异提示）的 AI prompt"""
    prompt = (
        "[RULE_TYPE:14]\n"
        "你是一个功能点估算审核专家助手。"
        "请判断以下每一对功能名称是否仅因单个名词不同而产生歧义，"
        "是否存在重复估算或描述过于相近的风险。\n\n"
        "判断标准：\n"
        "1. 如果两个名称的结构完全一致，仅一个宾语/对象名词不同"
        "（如\"新增用户\"和\"新增角色\"），且这两个名词在业务上有明显区别、"
        "代表不同的业务实体，则 aiConfirmed=false（不构成问题，可视为独立功能）。\n"
        "2. 如果仅一个名词不同，但该名词是同义词或近义词"
        "（如\"信息\"和\"资料\"、\"数据\"和\"记录\"），"
        "或者该名词在业务语境下含义重叠，则 aiConfirmed=true（需要复核）。\n"
        "3. 如果两个名称的差异不止一处（动词也不同），"
        "请标记 aiConfirmed=false（不在此规则检查范围内）。\n"
        "4. 注意区分享有同一组动词-宾语结构的CRUD操作"
        "（如\"新增用户\"/\"修改用户\"/\"删除用户\"/\"查询用户\"），"
        "这些属于不同的功能点，aiConfirmed=false。\n\n"
        "请按以下JSON格式返回结果，不要包含其他内容：\n"
        '[\n'
        '  {"nameI": "名称1", "nameJ": "名称2", "aiConfirmed": true/false, "reason": "判断理由"}\n'
        ']\n\n'
        "检查项列表：\n"
    )
    for idx, it in enumerate(items, 1):
        prompt += (
            f"\n{idx}. 功能名称A: {it['nameI']}\n"
            f"   功能名称B: {it['nameJ']}\n"
            f"   所在Sheet: {it.get('context', {}).get('sheet', '')}\n"
            f"   分组键(A-E列): {it.get('context', {}).get('aeGroup', '')}\n"
        )
    return prompt


@app.route("/api/ai/batch-match", methods=["POST"])
def ai_batch_match():
    raw = request.get_json(force=True)

    ai_url_input = ""
    rule_type = "2"
    items_input = raw

    if isinstance(raw, dict):
        ai_url_input = raw.get("aiUrl", "").strip()
        rule_type = str(raw.get("ruleType", "2")).strip() or "2"
        items_input = raw.get("items", [])

    if not isinstance(items_input, list) or len(items_input) == 0:
        return jsonify({"error": "请求参数应为非空数组"}), 400

    if AI_MOCK_MODE:
        base_url = ai_url_input or "mock://ai"
        print(f"[AI] Mock 模式: base_url={base_url}, ruleType={rule_type}")
    elif ai_url_input:
        base_url = ai_url_input
        print(f"[AI] 使用前端传入的 URL: {base_url}, ruleType={rule_type}")
    else:
        base_url = get_available_ai_base_url()
        if not base_url:
            return jsonify({"error": "AI 大模型接口不可用", "results": []}), 503

    # 根据 ruleType 构建不同的 items 和 prompt
    if rule_type == "8":
        items = _build_rule8_items(items_input)
        if not items:
            return jsonify({"error": "没有有效的规则8检查项", "results": []}), 400
        prompt = _build_rule8_prompt(items)
    elif rule_type == "14":
        items = _build_rule14_items(items_input)
        if not items:
            return jsonify({"error": "没有有效的规则14检查项", "results": []}), 400
        prompt = _build_rule14_prompt(items)
    else:
        # 规则2（默认，向后兼容）
        items = _build_rule2_items(items_input)
        if not items:
            return jsonify({"error": "没有有效的检查项", "results": []}), 400
        prompt = _build_rule2_prompt(items)

    print("\n" + "=" * 60)
    print(f"[AI Match] 规则{rule_type} 请求参数:")
    print(json.dumps(items, ensure_ascii=False, indent=2))
    print("=" * 60)

    # 打印发送给 AI 大模型的完整 prompt
    print("\n" + "=" * 80)
    print(f"[AI Match] 发送给 AI 大模型的完整 prompt (规则{rule_type}):")
    print(prompt)
    print("=" * 80 + "\n")

    if AI_MOCK_MODE:
        json_res = MockAIProvider.chat(base_url, "mock_session", prompt)
        print(f"\n[AI Match] Mock AI 接口返回原始内容:")
    else:
        session_id = ai_init_session(base_url)
        if not session_id:
            return jsonify({"error": "AI session 初始化失败", "results": []}), 500

        json_res = ai_chat(base_url, session_id, prompt)
        if not json_res:
            return jsonify({"error": "AI 调用返回为空", "results": []}), 500

        print(f"\n[AI Match] AI 接口返回原始内容:")

    print(json_res)
    print("-" * 60)

    # 通用解析：透传 AI 返回的所有字段，并附加 ruleType
    results = []
    try:
        parsed = json.loads(json_res) if isinstance(json_res, str) else json_res
        if isinstance(parsed, list):
            for r in parsed:
                item = {"ruleType": rule_type}
                item.update(r)
                results.append(item)
        elif isinstance(parsed, dict):
            item = {"ruleType": rule_type}
            item.update(parsed)
            results.append(item)
    except (json.JSONDecodeError, TypeError) as e:
        print(f"[AI Match] JSON 解析失败: {e}")

    print(f"\n[AI Match] 解析后的结果:")
    print(json.dumps(results, ensure_ascii=False, indent=2))
    print("=" * 60)

    return jsonify({"results": results})


def _build_rule2_items(items_input: list[dict]) -> list[dict]:
    """构建规则2的检查项列表，自动补全 wpsDesc"""
    items = []
    for item in items_input:
        batch = item.get("batch", "")
        point_index = item.get("pointIndex", "")
        point_name = item.get("pointName", "")
        wps_desc = item.get("wpsDesc", "")
        if not batch or not point_index or not point_name:
            continue

        if not wps_desc and batch in _batch_docs:
            cm = _batch_docs[batch].get("content_map", {})
            if point_index in cm:
                wps_desc = cm[point_index].get("content", "")

        items.append({
            "batch": batch,
            "pointIndex": point_index,
            "pointName": point_name,
            "wpsDesc": wps_desc,
        })
    return items


def _build_rule2_prompt(items: list[dict]) -> str:
    """构造规则2（需求依据校验）的 AI prompt"""
    prompt = (
        "[RULE_TYPE:2]\n"
        "你是一个需求文档检查助手。请判断以下每个功能名称是否在其对应的需求文档章节内容中有依据。\n"
        "请按以下JSON格式返回结果，不要包含其他内容：\n"
        '[\n'
        '  {"batch": "批次名", "pointIndex": "章节号", "pointName": "功能名", "isInWPS": true/false, "reason": "判断理由"}\n'
        ']\n\n'
        "检查项列表：\n"
    )
    for idx, it in enumerate(items, 1):
        prompt += (
            f"\n{idx}. 批次: {it['batch']}\n"
            f"   章节: {it['pointIndex']}\n"
            f"   功能名称: {it['pointName']}\n"
            f"   需求文档内容: {it['wpsDesc'][:500] if it['wpsDesc'] else '(无内容)'}\n"
        )
    return prompt


def _build_rule8_items(items_input: list[dict]) -> list[dict]:
    """构建规则8的检查项列表，验证必要字段"""
    items = []
    for item in items_input:
        name_i = (item.get("nameI") or "").strip()
        name_j = (item.get("nameJ") or "").strip()
        if not name_i or not name_j:
            continue
        items.append({
            "nameI": name_i,
            "nameJ": name_j,
            "rule8Kind": item.get("rule8Kind", "near"),
            "context": item.get("context", {}),
        })
    return items


def _build_rule14_items(items_input: list[dict]) -> list[dict]:
    """构建规则14的检查项列表，验证必要字段"""
    items = []
    for item in items_input:
        name_i = (item.get("nameI") or "").strip()
        name_j = (item.get("nameJ") or "").strip()
        if not name_i or not name_j:
            continue
        items.append({
            "nameI": name_i,
            "nameJ": name_j,
            "context": item.get("context", {}),
        })
    return items


@app.route("/api/debug/upload", methods=["POST"])
def debug_upload():
    if "file" not in request.files:
        return jsonify({"error": "缺少文件"}), 400

    file = request.files["file"]
    batch = request.form.get("batch", "")
    safe_name = "".join(c if c.isalnum() or c in "._- " else "_" for c in file.filename)
    temp_dir = _get_temp_session_dir()
    temp_path = os.path.join(temp_dir, f"debug_{safe_name}")
    file.save(temp_path)

    file_size = os.path.getsize(temp_path)
    with open(temp_path, "rb") as f:
        header = f.read(8)
    is_ole2 = header[:8] == b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"
    is_zip = header[:2] == b"PK"

    text = ""
    extract_error = ""
    extract_method = ""
    table_count = 0
    heading_stats = {}
    try:
        from docx import Document
        doc = Document(temp_path)
        table_count = len(doc.tables)
        for p in doc.paragraphs:
            if p.style:
                sn = (p.style.name or '').lower()
                if 'heading' in sn:
                    heading_stats[sn] = heading_stats.get(sn, '') or ''
        text = parser._extract_from_docx(doc)
        extract_method = "python-docx-enhanced"
    except Exception as e:
        extract_method = "fallback"
        try:
            text = parser.extract_text_fallback(temp_path)
        except Exception as e2:
            extract_error = str(e2)

    toc = parser.parse_toc_structure(text) if text else {}
    content_map = parser.extract_all_subsections_with_content(text, toc) if toc else {}

    safe_filename = "".join(c if c.isalnum() or c in "._- " else "_" for c in file.filename)
    dump_path = ""
    try:
        dump_path = os.path.join(temp_dir, f"debug_rawtext_{safe_filename}.txt")
        with open(dump_path, "w", encoding="utf-8") as f:
            f.write(text)
    except Exception as e:
        print(f"[DEBUG] 保存原始文本失败: {e}")
        dump_path = ""

    print("\n" + "=" * 80)
    print(f"【调试模式 - 批次: {batch} - 文件: {file.filename}】")
    print(f"  文件大小: {file_size} bytes")
    print(f"  OLE2 格式: {is_ole2}")
    print(f"  ZIP/doxc 格式: {is_zip}")
    print(f"  提取方式: {extract_method}")
    if extract_error:
        print(f"  提取错误: {extract_error}")
    print(f"  全文长度: {len(text)} 字符")
    print(f"  表格数量: {table_count}")
    print(f"  标题层级: {list(heading_stats.keys())}")
    print(f"  全文已保存至: {dump_path}")
    print(f"\n{'─' * 80}")
    print("【原始全文（开头 2000 字符）】")
    print(text[:2000])
    print(f"\n{'─' * 80}")
    print("【原始全文（末尾 2000 字符）】")
    print(text[-2000:] if len(text) > 2000 else "(全文不足2000字符)")
    print(f"\n{'─' * 80}")
    print(f"【TOC 解析结果】（共 {len(toc)} 项）")
    for key in sorted(toc.keys(), key=lambda k: (int(k.split(".")[0]), float(k) if "." in k else 0)):
        print(f"  [{key}] {toc[key]}")
    print(f"\n【正文提取结果】（共 {len(content_map)} 个子节）")
    for key in sorted(content_map.keys(), key=lambda k: float(k)):
        item = content_map[key]
        content_len = len(item["content"]) if item["content"] else 0
        preview = item["content"][:100] if item["content"] else "(无内容)"
        print(f"  [{key}] {item['title']}  ({content_len} chars)")
        print(f"      预览: {preview}")
    print("=" * 80 + "\n")

    try:
        os.remove(temp_path)
    except Exception:
        pass

    wps_list = []
    # 需要过滤的标题列表（完全匹配）
    ERROR_INFO_TITLES = {"错误信息", "错误信息(可选)", "错误信息（可选）"}

    def key_to_tuple(k):
        return tuple(int(x) for x in k.split('.'))
    for key in sorted(content_map.keys(), key=key_to_tuple):
        parts = key.split('.')
        if len(parts) >= 2 and parts[0] == '3' and int(parts[1]) >= 4:
            item = content_map[key]
            title = item.get("title", "")
            # 过滤掉"错误信息"相关小节（完全匹配）
            if title in ERROR_INFO_TITLES:
                continue
            wps_list.append({
                "batch": batch,
                "index": key,
                "wpsDesc": item.get("title", "") + item.get("content", ""),
            })

    return jsonify({
        "batch": batch,
        "fileName": file.filename,
        "fileSize": file_size,
        "isOLE2": is_ole2,
        "isZip": is_zip,
        "extractMethod": extract_method,
        "extractError": extract_error,
        "textLength": len(text),
        "tableCount": table_count,
        "headingStyles": list(heading_stats.keys()),
        "textHead500": text[:500],
        "textTail500": text[-500:] if len(text) > 500 else "",
        "textHead2000": text[:2000],
        "textTail2000": text[-2000:] if len(text) > 2000 else "",
        "tocCount": len(toc),
        "toc": toc,
        "contentMapCount": len(content_map),
        "contentMap": {k: {"title": v["title"], "contentLen": len(v.get("content", "")), "contentPreview": v.get("content", "")[:200]} for k, v in content_map.items()},
        "tocConsistency": check_toc_body_consistency(text, toc),
        "wpsList": wps_list,
        "dumpPath": dump_path,
    })


@app.route("/agentInitSession", methods=["POST"])
def agent_init_session():
    data = request.get_json(force=True)
    base_url = (data.get("baseUrl") or "").rstrip("/")
    in_both = data.get("inBoth") or []
    rule8_items = data.get("rule8Items") or []
    rule14_items = data.get("rule14Items") or []

    if AI_MOCK_MODE:
        import uuid
        session_id = "mock-" + uuid.uuid4().hex[:12]
        _mock_session_meta[session_id] = {
            "checkListLen": len(in_both),
            "rule8Items": rule8_items,
            "rule14Items": rule14_items,
        }
        print(f"[Agent][Mock] 模拟会话初始化 - session_id: {session_id} "
              f"(checkList={len(in_both)}, rule8={len(rule8_items)}, rule14={len(rule14_items)})")
        return jsonify({"data": {"session_id": session_id}})

    if not base_url:
        return jsonify({"error": "缺少 baseUrl"}), 400

    # payload 使用 config_variables 数组格式，兼容远程 AI 服务
    # rule8/rule14 只传 nameI、nameJ，去掉其他属性
    rule8_for_ai = [{"nameI": it.get("nameI", ""), "nameJ": it.get("nameJ", "")} for it in rule8_items]
    rule14_for_ai = [{"nameI": it.get("nameI", ""), "nameJ": it.get("nameJ", "")} for it in rule14_items]

    url = base_url + "/chatabc/init_session"
    payload = {
        "appId": "",
        "trCode": "",
        "trVersion": "",
        "timestamp": "1",
        "requestId": "",
        "data": {
            "config_variables": [
                {"name": "checkList", "value": json.dumps(in_both, ensure_ascii=False)},
                {"name": "resList", "value": "[]"},
                {"name": "rule8Items", "value": json.dumps(rule8_for_ai, ensure_ascii=False)},
                {"name": "resListRule8", "value": "[]"},
                {"name": "rule14Items", "value": json.dumps(rule14_for_ai, ensure_ascii=False)},
                {"name": "resListRule14", "value": "[]"},
            ]
        },
    }

    print(f"[Agent] 会话初始化 - baseUrl: {base_url} "
          f"(checkList={len(in_both)}, rule8={len(rule8_items)}, rule14={len(rule14_items)})")

    # 打印发送给 AI 大模型的完整报文
    print("\n" + "=" * 80)
    print("[Agent] 发送给 AI 大模型的 init_session 报文:")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    print("=" * 80 + "\n")

    try:
        resp = requests.post(url, json=payload, timeout=60)
        result = resp.json()
        session_id = result.get("data", {}).get("session_id", "")
        if not session_id:
            print("[Agent] init_session 未返回 session_id")
            return jsonify({"error": "init_session 未返回 session_id"}), 500
        print(f"[Agent] session_id 获取: {session_id}")
        return jsonify(result)
    except Exception as e:
        print(f"[Agent] 会话初始化失败: {e}")
        return jsonify({"error": f"会话初始化失败: {e}"}), 500


@app.route("/agentInitChat", methods=["POST"])
def agent_init_chat():
    data = request.get_json(force=True)
    base_url = (data.get("baseUrl") or "").rstrip("/")
    session_id = data.get("sessionId") or ""

    if AI_MOCK_MODE:
        print(f"[Agent][Mock] 模拟对话 - session_id: {session_id}")

        # 从 session 元数据获取各分支需要生成的数量和实际 items
        meta = _mock_session_meta.pop(session_id, {})
        check_len = meta.get("checkListLen", 0)
        rule8_items = meta.get("rule8Items", [])
        rule14_items = meta.get("rule14Items", [])

        # 生成三合一格式的 mock 结果（使用实际的功能名称）
        check_list_result = MockAIProvider.mock_checklist_result(check_len)
        rule8_result = MockAIProvider.mock_rule8_result(len(rule8_items), rule8_items)
        rule14_result = MockAIProvider.mock_rule14_result(len(rule14_items), rule14_items)

        print(f"[Agent][Mock] 生成结果 - checkList: {check_len}条, rule8: {len(rule8_result)}条, rule14: {len(rule14_result)}条")
        print(f"[Agent][Mock] rule8Result: {json.dumps(rule8_result, ensure_ascii=False)}")
        print(f"[Agent][Mock] rule14Result: {json.dumps(rule14_result, ensure_ascii=False)}")

        mock_output = json.dumps({
            "checkListRes": check_list_result,
            "rule8Result": rule8_result,
            "rule14Result": rule14_result,
        }, ensure_ascii=False)

        mock_event = json.dumps({
            "additional_kwargs": {
                "node_id": "end",
                "node_output": {"output": mock_output},
            }
        }, ensure_ascii=False)

        def mock_generate():
            import random as _rand
            import time as _time
            # 模拟 AI 工作流节点执行过程（含 type 字段 + node_title）
            thinking_steps = [
                {"node_id": "node_check_1", "node_title": "需求文档解析节点", "type": "thinking",
                 "node_output": {"output": "正在分析功能点与需求文档的对应关系..."}},
                {"node_id": "node_check_2", "node_title": "批次估计书比对节点", "type": "thinking",
                 "node_output": {"output": "正在逐项比对批次估计书内容..."}},
                {"node_id": "node_check_3", "node_title": "语义相似度计算节点", "type": "thinking",
                 "node_output": {"output": "正在计算语义相似度..."}},
                {"node_id": "node_rule8_1", "node_title": "规则8交叉验证节点", "type": "analyzing",
                 "node_output": {"output": "正在交叉验证规则8的重复项..."}},
                {"node_id": "node_rule14_1", "node_title": "单名词差异匹配节点", "type": "thinking",
                 "node_output": {"output": "正在匹配单名词差异（规则14）..."}},
                {"node_id": "node_summary_1", "node_title": "分析结果汇总节点", "type": "analyzing",
                 "node_output": {"output": "正在汇总分析结果..."}},
            ]

            # 随机选取 0~2 个节点注入异常（用于前端测试）
            import copy
            error_candidates = [
                {"node_id": "node_rule8_1", "node_title": "规则8交叉验证节点", "type": "ai",
                 "node_output": {"output": {"exception": "JSON 解析错误：AI 返回的规则8验证结果格式不符合预期 schema，缺少必填字段 aiConfirmed"}}},
                {"node_id": "node_check_3", "node_title": "语义相似度计算节点", "type": "ai",
                 "node_output": {"output": {"exception": "连接超时：调用语义相似度计算服务超时（30s），请检查下游服务是否正常运行"}}},
                {"node_id": "node_summary_1", "node_title": "分析结果汇总节点", "type": "ai",
                 "node_output": {"output": {"exception": "部分数据丢失：规则14的输入数据在传输过程中损坏，节点接收到空列表"}}},
            ]
            num_errors = _rand.randint(0, 2)
            error_events = _rand.sample(error_candidates, num_errors) if num_errors > 0 else []

            for step in thinking_steps:
                _time.sleep(_rand.uniform(0.4, 1.0))
                event_data = json.dumps({
                    "content": "",
                    "additional_kwargs": {
                        "node_id": step["node_id"],
                        "node_title": step["node_title"],
                        "node_output": step["node_output"],
                    },
                    "response_metadata": {},
                    "type": step.get("type", "thinking"),
                }, ensure_ascii=False)
                yield f"data: {event_data}\n\n"

            # 注入异常事件（在最终结果之前，模拟工作流中节点报错）
            for err_evt in error_events:
                _time.sleep(_rand.uniform(0.3, 0.8))
                err_data = json.dumps({
                    "content": "",
                    "additional_kwargs": {
                        "node_id": err_evt["node_id"],
                        "node_title": err_evt["node_title"],
                        "node_output": err_evt["node_output"],
                    },
                    "response_metadata": {},
                    "type": "ai",
                }, ensure_ascii=False)
                yield f"data: {err_data}\n\n"

            # 最终结果
            _time.sleep(_rand.uniform(0.8, 1.8))
            yield f"data: {mock_event}\n\n"
            yield "data: [DONE]\n\n"

        return Response(
            stream_with_context(mock_generate()),
            mimetype="text/event-stream",
        )

    if not base_url:
        return jsonify({"error": "缺少 baseUrl"}), 400
    if not session_id:
        return jsonify({"error": "缺少 sessionId"}), 400

    url = base_url + "/chatabc/chat"
    payload = {
        "appId": "",
        "trCode": "",
        "trVersion": "",
        "timestamp": "1",
        "requestId": "",
        "data": {
            "session_id": session_id,
            "txt": "123",
            "files": [],
            "stream": True,
        },
    }

    # 打印发送给 AI 大模型的完整报文
    print("\n" + "=" * 80)
    print(f"[Agent] 发送给 AI 大模型的 chat 报文 (session_id: {session_id}):")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    print("=" * 80 + "\n")

    try:
        def generate():
            # timeout=(连接超时, 读取超时) - 连接超时 30s，读取超时 5 分钟
            remote_resp = requests.post(url, json=payload, stream=True, timeout=(30, 300))
            line_count = 0
            for line in remote_resp.iter_lines(decode_unicode=True):
                if line:
                    line_count += 1
                    yield f"data: {line}\n\n"
            yield f"data: [DONE]\n\n"
            print(f"[Agent] 会话对话完成 - session_id: {session_id} - 共 {line_count} 行")

        return Response(
            stream_with_context(generate()),
            mimetype="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "X-Accel-Buffering": "no",
            },
        )
    except Exception as e:
        print(f"[Agent] 会话对话失败: {e}")
        return jsonify({"error": f"对话失败: {e}"}), 500


def _auto_cleanup_loop():
    """后台守护线程：每隔 CLEANUP_INTERVAL_SECONDS 自动清理过期临时文件"""
    print(f"[自动清理] 后台清理线程已启动（间隔={CLEANUP_INTERVAL_SECONDS}s, 过期阈值={CLEANUP_EXPIRY_SECONDS}s）")
    while True:
        time.sleep(CLEANUP_INTERVAL_SECONDS)
        try:
            _do_cleanup_temp_files()
        except Exception as e:
            print(f"[自动清理] 清理过程出错: {e}")


# ═══════════════════════════════════════════════════════════
#  版本检查与自动更新 API
# ═══════════════════════════════════════════════════════════

# 内网 Git 平台认证 Cookie（硬编码，修改后需重启服务）
COOKIE = ""

def _get_config() -> dict:
    """获取更新配置，注入 Cookie"""
    config = version_manager.read_update_config()
    if COOKIE:
        config["cookie"] = COOKIE
    return config


@app.route("/api/version/current", methods=["GET"])
def api_version_current():
    """获取当前版本信息"""
    try:
        local = version_manager.read_local_version()
        config = _get_config()
        return jsonify({
            "version": local.get("version", "0.0.0.0"),
            "update_enabled": config.get("enabled", False),
            "auto_check": config.get("auto_check", False),
        })
    except Exception as e:
        return jsonify({"error": f"读取版本信息失败: {str(e)}"}), 500


@app.route("/api/version/check", methods=["GET"])
def api_version_check():
    """检查是否有新版本（轻量操作，仅下载 version.json）"""
    try:
        config = _get_config()
        if not config.get("enabled", False):
            return jsonify({
                "update_available": False,
                "current_version": "0.0.0.0",
                "latest_version": None,
                "changelog": [],
                "package_size": 0,
                "error": "更新功能未启用",
            })

        result = version_manager.check_for_update(config)
        return jsonify(result)
    except Exception as e:
        return jsonify({
            "update_available": False,
            "error": f"版本检查失败: {str(e)}",
        }), 500


@app.route("/api/version/update", methods=["POST"])
def api_version_update():
    """触发更新流程（后台异步执行）"""
    try:
        config = _get_config()
        if not config.get("enabled", False):
            return jsonify({"status": "error", "message": "更新功能未启用"}), 400

        result = version_manager.do_update(config)
        return jsonify(result)
    except Exception as e:
        return jsonify({"status": "error", "message": f"启动更新失败: {str(e)}"}), 500


@app.route("/api/version/status", methods=["GET"])
def api_version_status():
    """查询更新进度"""
    try:
        state = version_manager.get_update_state()
        return jsonify({
            "phase": state.get("phase", "idle"),
            "progress_percent": state.get("progress_percent", 0),
            "stage_message": state.get("stage_message", ""),
            "detail_message": state.get("detail_message", ""),
            "error_code": state.get("error_code"),
            "error_message": state.get("error_message", ""),
            "can_retry": state.get("can_retry", False),
            "last_check_time": state.get("last_check_time"),
            "last_update_time": state.get("last_update_time"),
        })
    except Exception as e:
        return jsonify({"error": f"查询状态失败: {str(e)}"}), 500


@app.route("/api/version/rollback", methods=["POST"])
def api_version_rollback():
    """回滚到指定备份"""
    try:
        data = request.get_json(silent=True) or {}
        backup_id = data.get("backup_id", "")

        if not backup_id:
            # 未指定备份 ID，列出所有可用备份
            backups = version_manager.list_backups()
            return jsonify({
                "backups": backups,
                "message": f"共 {len(backups)} 个可用备份" if backups else "没有可用的备份",
            })

        result = version_manager.rollback_to_backup(backup_id)
        status_code = 200 if result.get("status") == "rolled_back" else 400
        return jsonify(result), status_code
    except Exception as e:
        return jsonify({"status": "error", "message": f"回滚失败: {str(e)}"}), 500


@app.route("/api/version/debug", methods=["GET"])
def api_version_debug():
    """调试端点：分步展示版本检查的每个环节，帮助排查配置问题"""
    try:
        config = _get_config()
        steps = []

        # Step 1: 配置检查
        steps.append({
            "step": 1, "name": "读取配置",
            "enabled": config.get("enabled", False),
            "releases_api_url": config.get("releases_api_url", ""),
            "cookie_file": config.get("cookie_file", ""),
        })

        # Step 2: Cookie 加载
        cookie_path = os.path.join(version_manager._get_base_dir(),
                                   config.get("cookie_file", "cookie.txt"))
        cookie_exists = os.path.exists(cookie_path)
        cookie_raw = ""
        cookie_decoded = ""
        cookie_error = None

        if cookie_exists:
            try:
                with open(cookie_path, 'r', encoding='utf-8') as f:
                    cookie_raw = f.read().strip()
                lines = [l.strip() for l in cookie_raw.split('\n')
                         if l.strip() and not l.strip().startswith('#')]
                if lines:
                    import base64
                    try:
                        cookie_decoded = base64.b64decode(lines[0]).decode('utf-8')
                    except Exception as e:
                        cookie_error = f"Base64 解码失败: {e}"
                else:
                    cookie_error = "文件中没有有效的 Cookie 内容（全是空行或注释）"
            except IOError as e:
                cookie_error = f"读取文件失败: {e}"

        steps.append({
            "step": 2, "name": "Cookie 加载",
            "cookie_path": cookie_path,
            "cookie_exists": cookie_exists,
            "cookie_raw_length": len(cookie_raw),
            "cookie_raw_preview": cookie_raw[:60] + "..." if len(cookie_raw) > 60 else cookie_raw,
            "cookie_decoded_length": len(cookie_decoded),
            "cookie_decoded_preview": cookie_decoded[:60] + "..." if len(cookie_decoded) > 60 else cookie_decoded,
            "error": cookie_error,
        })

        if not config.get("enabled", False):
            return jsonify({"steps": steps, "summary": "更新功能未启用"})

        # Step 3: 本地版本
        local = version_manager.read_local_version()
        steps.append({
            "step": 3, "name": "本地版本",
            "version": local.get("version", "0.0.0.0"),
        })

        # Step 4: HTTP 请求
        releases_url = config.get("releases_api_url", "")
        headers = version_manager._build_auth_headers(config)
        has_cookie = "Cookie" in headers

        http_error = None
        http_status = None
        html_length = 0
        html_preview = ""

        if not releases_url:
            http_error = "releases_api_url 未配置"
        else:
            try:
                resp = version_manager._http_get(releases_url, config, timeout=30)
                if resp is None:
                    http_error = "HTTP 请求失败（连接被拒绝或超时）"
                else:
                    http_status = resp.status_code
                    html_text = resp.text or ""
                    html_length = len(html_text)
                    html_preview = html_text[:300]
                    if resp.status_code == 401 or resp.status_code == 403:
                        http_error = f"认证失败 (HTTP {resp.status_code})，Cookie 无效或已过期"
                    elif resp.status_code != 200:
                        http_error = f"服务器返回异常状态码: HTTP {resp.status_code}"
            except Exception as e:
                http_error = f"请求异常: {e}"

        steps.append({
            "step": 4, "name": "HTTP 请求",
            "url": releases_url,
            "has_cookie": has_cookie,
            "cookie_header_preview": headers.get("Cookie", "")[:60] + "..."
                if headers.get("Cookie", "") else "",
            "http_status": http_status,
            "html_length": html_length,
            "html_preview": html_preview,
            "error": http_error,
        })

        # Step 5: HTML 解析
        parse_error = None
        links_count = 0
        links = []
        if http_status == 200 and html_length > 0:
            try:
                releases = version_manager.fetch_releases(config)
                if releases is None:
                    parse_error = "fetch_releases 返回 None（HTML 解析失败或无下载链接）"
                else:
                    links_count = len(releases)
                    links = [{"version": r["version"], "url": r["url"][:100]}
                             for r in releases[:5]]
            except Exception as e:
                parse_error = f"HTML 解析异常: {e}"

        steps.append({
            "step": 5, "name": "HTML 解析",
            "links_count": links_count,
            "links": links,
            "error": parse_error,
        })

        # 汇总
        summary = "就绪" if not http_error and not parse_error and links_count > 0 else "存在问题"
        return jsonify({"steps": steps, "summary": summary})

    except Exception as e:
        return jsonify({"error": f"调试失败: {str(e)}"}), 500


def _port_in_use(port: int) -> bool:
    """检查端口是否已被占用"""
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            s.settimeout(1)
            return s.connect_ex(("127.0.0.1", port)) == 0
    except Exception:
        return False


def _startup_version_check():
    """启动后自动检查版本更新（后台线程）"""
    import time
    time.sleep(3)  # 等待 Flask 服务完全启动
    config = version_manager.read_update_config()
    if config.get("enabled") and config.get("auto_check"):
        try:
            result = version_manager.check_for_update()
            if result.get("update_available"):
                print(f"[Startup] 发现新版本: {result.get('latest_version')}")
        except Exception as e:
            print(f"[Startup] 自动版本检查失败: {e}")


# ═══════════════════════════════════════════════════════════
#  Native Messaging 模式（合并原 start_backend_host，单 exe 双角色）
#  Chrome/Edge 以 NMH 方式启动本 exe 时，argv 会携带 chrome-extension:// 来源；
#  此时仅负责"查端口 → 分离重启自身（服务模式）→ 应答 → 退出"。
# ═══════════════════════════════════════════════════════════

def _is_native_messaging_invocation() -> bool:
    return any("chrome-extension://" in arg for arg in sys.argv[1:])


def _run_native_messaging_host():
    import struct

    def _read_msg():
        raw = sys.stdin.buffer.read(4)
        if len(raw) < 4:
            return None
        length = struct.unpack("I", raw)[0]
        return json.loads(sys.stdin.buffer.read(length).decode("utf-8"))

    def _write_msg(obj):
        data = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        sys.stdout.buffer.write(struct.pack("I", len(data)))
        sys.stdout.buffer.write(data)
        sys.stdout.buffer.flush()

    try:
        _read_msg()
    except Exception:
        return

    if _port_in_use(8765):
        try:
            _write_msg({"ok": True, "started": False, "reason": "already_running"})
        except Exception:
            pass
        return

    # 以服务模式分离重启自身（PyInstaller 下 sys.executable 即 exe 自身；源码模式需带脚本路径）
    if getattr(sys, "frozen", False):
        cmd = [sys.executable]
    else:
        cmd = [sys.executable, os.path.abspath(__file__)]
    flags = getattr(subprocess, "CREATE_NEW_CONSOLE", 0) if os.name == "nt" else 0
    try:
        subprocess.Popen(cmd, creationflags=flags, close_fds=True,
                         stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                         stderr=subprocess.DEVNULL)
        _write_msg({"ok": True, "started": True})
        print("[NMH] 已分离启动后端服务进程")
    except Exception as e:
        try:
            _write_msg({"ok": False, "reason": str(e)})
        except Exception:
            pass


if __name__ == "__main__":
    # NMH 模式：应答浏览器后退出，不启动第二个服务实例
    if _is_native_messaging_invocation():
        _run_native_messaging_host()
        sys.exit(0)

    # 检查是否已有实例在运行，避免 projecttool:// 重复拉起
    if _port_in_use(8765):
        print("[Startup] 端口 8765 已被占用，已有实例在运行，退出。")
        sys.exit(0)

    # 确保 PyInstaller exe 运行目录有可写的配置文件
    version_manager.ensure_runtime_configs()

    # 启动自动清理后台线程（守护线程，程序退出时自动终止）
    cleanup_thread = threading.Thread(target=_auto_cleanup_loop, daemon=True, name="auto-cleanup")
    cleanup_thread.start()

    # 启动后自动检查版本更新（后台线程，不阻塞服务启动）
    update_check_thread = threading.Thread(
        target=_startup_version_check, daemon=True, name="startup-update-check"
    )
    update_check_thread.start()

    app.run(host="127.0.0.1", port=8765, debug=False)
